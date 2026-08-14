"""Étape 4 — orchestration de l'enrichissement nominatif (Paris, tous lycées).

Usage :
    python -m ie_prospection.enrich_nominatif.pipeline4
    python -m ie_prospection.enrich_nominatif.pipeline4 --enable-linkedin  # si accès légal

Traite TOUS les lycées du fichier Paris (aucun filtre P1/P2/P3). Ne contacte
personne, n'importe rien dans HubSpot, n'invente aucune donnée.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from ..logging_utils import get_logger, setup_logging
from .build4 import build_nominatif
from .robots import RobotsPolicy
from .schema4 import CONTACT_COLUMNS, ETAB_COLUMNS
from .sources4 import NominatifConfig, gather_for_etab

INPUT_CANDIDATES = [
    "data/output/paris/03_base_finale_paris.xlsx",   # prioritaire (étape 3)
    "data/output/paris/02_enrichissement_public_paris.xlsx",
]


def load_input(path: str | None) -> tuple[list[dict], str]:
    candidates = [path] if path else INPUT_CANDIDATES
    for cand in candidates:
        if cand and Path(cand).exists():
            return _load_xlsx(cand), cand
    raise FileNotFoundError("Aucun fichier d'entrée trouvé (cherché : "
                            + ", ".join(c for c in candidates if c) + ")")


def _load_xlsx(path: str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        out.append({header[i]: ("" if v is None else str(v)) for i, v in enumerate(r)})
    return out


def _write_xlsx_multi(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    first = True
    for name, (cols, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name[:31]
        first = False
        ws.append(cols)
        for row in rows:
            ws.append([_txt(row.get(c, "")) for c in cols])
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(c) + 2, 12), 42)
        ws.freeze_panes = "A2"
    wb.save(path)
    return path


def _txt(v):
    return "" if v is None else str(v)


def _write_csv(path: Path, rows: list[dict], columns: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    return path


def run(input_path: str | None, out_dir: str, reports_dir: str, run_id: str,
        cfg: NominatifConfig, session=None) -> dict:
    log = get_logger()
    date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    etab_rows, used = load_input(input_path)
    log.info("Étape 4 — entrée : %s (%d lycées, AUCUN filtre P1/P2/P3)", used, len(etab_rows))

    session = session or requests.Session()
    policy = RobotsPolicy(user_agent=cfg.user_agent, timeout=cfg.timeout, session=session)

    results = []
    for i, etab in enumerate(etab_rows, 1):
        contacts, attempts = gather_for_etab(etab, cfg, policy, session, date)
        results.append((etab, contacts, attempts))
        if i % 25 == 0:
            log.info("… %d/%d lycées traités", i, len(etab_rows))

    built = build_nominatif(results, date)

    out_dir, reports_dir = Path(out_dir), Path(reports_dir)
    xlsx_path = _write_xlsx_multi(
        out_dir / "04_enrichissement_nominatif_paris.xlsx",
        {"Établissements": (ETAB_COLUMNS, built["etablissements"]),
         "Contacts": (CONTACT_COLUMNS, built["contacts"])})

    r1 = _write_csv(reports_dir / "04_lycees_sans_contact.csv", built["sans_contact"],
                    ["UAI", "Nom", "Académie", "Site web", "Statut", "Raisons",
                     "Date dernière recherche"])
    r2 = _write_csv(reports_dir / "04_emails_a_verifier.csv", built["emails_a_verifier"],
                    ["UAI établissement", "Prénom", "Nom", "Rôle", "Email",
                     "Type email", "Email vérifié O/N", "Source de la donnée",
                     "URL source", "Date de collecte"])
    r3 = _write_csv(reports_dir / "04_sources_utilisees.csv", built["sources_report"],
                    ["étape", "source", "statut", "occurrences"])

    s = built["stats"]
    log.info("=" * 60)
    log.info("Étape 4 : lycées=%d | contacts=%d | sans contact=%d | emails à vérifier=%d",
             s["etablissements"], s["contacts"], s["lycees_sans_contact"],
             s["emails_a_verifier"])
    log.info("Fichiers : %s | %s | %s | %s", xlsx_path, r1, r2, r3)
    return {"stats": s, "files": {"xlsx": str(xlsx_path), "sans_contact": str(r1),
                                  "emails_a_verifier": str(r2), "sources": str(r3)},
            "input": used}


def parse_args(argv=None):
    p = argparse.ArgumentParser(description="NEXA — étape 4 enrichissement nominatif (Paris)")
    p.add_argument("--input", default=None, help="Fichier d'entrée (xlsx)")
    p.add_argument("--out-dir", default="data/output/paris")
    p.add_argument("--reports-dir", default="reports/paris")
    p.add_argument("--user-agent", default="NEXA-IE-prospection-bot")
    p.add_argument("--timeout", type=float, default=8.0)
    p.add_argument("--max-pages-per-site", type=int, default=4)
    p.add_argument("--no-reconstruct-emails", action="store_true",
                   help="Désactiver la reconstruction d'emails académiques (marqués Non vérifié)")
    p.add_argument("--enable-linkedin", action="store_true",
                   help="N'activer QUE si un accès LinkedIn/SalesNav légal est déjà configuré")
    p.add_argument("--run-id", default=None)
    return p.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    run_id = args.run_id or datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    setup_logging(args.reports_dir, f"04_{run_id}")
    cfg = NominatifConfig(user_agent=args.user_agent, timeout=args.timeout,
                          max_pages_per_site=args.max_pages_per_site,
                          reconstruct_emails=not args.no_reconstruct_emails,
                          enable_linkedin=args.enable_linkedin)
    run(args.input, args.out_dir, args.reports_dir, run_id, cfg)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
