"""Écriture de fichiers .xlsx (openpyxl)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def write_xlsx(path: str | Path, rows: list[dict], columns: list[str],
               sheet_name: str = "Data") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append(columns)
    for row in rows:
        ws.append([_cell(row.get(c, "")) for c in columns])
    # Largeur de colonnes indicative + entête figée.
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(col) + 2, 12), 40)
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def _cell(value):
    # Les identifiants (UAI/SIREN/SIRET/CP) restent en TEXTE : on force str pour
    # préserver les zéros initiaux et éviter toute conversion numérique d'Excel.
    if value is None:
        return ""
    return str(value)
