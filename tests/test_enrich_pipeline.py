"""Intégration étape 2 : run complet fixture sur un socle Paris temporaire."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from ie_prospection.enrich.pipeline2 import FULL_COLUMNS, run
from ie_prospection.logging_utils import setup_logging
from ie_prospection.schema import OUTPUT_COLUMNS


def _make_socle(path: Path, n=120):
    rows = []
    for i in range(n):
        r = {c: "" for c in OUTPUT_COLUMNS}
        r["UAI"] = f"075{i:04d}A"
        r["Nom"] = f"Lycée {i}"
        r["Académie"] = "Paris"
        r["Campus rattaché"] = "Paris"
        r["Distance campus km"] = "12.0"
        r["Site web"] = f"https://l{i}.fr" if i % 2 == 0 else ""
        r["Date extraction"] = "2026-08-14T00:00:00Z"
        rows.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=OUTPUT_COLUMNS)
        w.writeheader()
        w.writerows(rows)


def test_enrich_end_to_end(tmp_path):
    setup_logging(tmp_path / "reports", "t2")
    socle = tmp_path / "in" / "campus_paris.csv"
    _make_socle(socle)
    out_dir = tmp_path / "out"
    reports_dir = tmp_path / "rep"

    res = run("fixture", str(socle), "config/campuses.yml", str(out_dir),
              str(reports_dir), radius=60.0, run_id="t2")

    # Tous les livrables existent.
    assert (out_dir / "02_enrichissement_public_paris.csv").exists()
    assert (out_dir / "02_enrichissement_public_paris.xlsx").exists()
    assert (reports_dir / "02_jointures.md").exists()
    assert (reports_dir / "02_anomalies.csv").exists()
    assert (reports_dir / "02_controle_30_lignes.xlsx").exists()

    # Le CSV enrichi a bien les colonnes socle + enrichissement.
    with (out_dir / "02_enrichissement_public_paris.csv").open(encoding="utf-8") as f:
        reader = csv.DictReader(f)
        assert reader.fieldnames == FULL_COLUMNS
        rows = list(reader)
    assert len(rows) == res["perimeter"] == 120

    # UAI conservé en texte dans le contrôle 30 lignes (zéros initiaux).
    wb = load_workbook(reports_dir / "02_controle_30_lignes.xlsx")
    ws = wb.active
    assert ws.max_row == 31  # 30 + entête
    assert isinstance(ws["A2"].value, str) and ws["A2"].value.startswith("075")

    # Au moins une source jointe et documentée.
    assert any(s["available"] and s["matched"] > 0 for s in res["join_stats"])
